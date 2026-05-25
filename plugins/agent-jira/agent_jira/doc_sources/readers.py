from __future__ import annotations

import sys
from pathlib import Path
from typing import List, Optional

from .ocr_backends import run_image_bytes_ocr, run_image_ocr, run_pdf_ocr
from .utils import normalize_text, strip_front_matter, warn_missing_dependency


def read_markdown(path: Path) -> Optional[str]:
    """Legge file Markdown rimuovendo eventuale front matter."""

    try:
        raw = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        try:
            raw = path.read_text(encoding="utf-8", errors="ignore")
        except OSError as exc:
            print(f"[filesystem] ⚠️ Errore leggendo {path}: {exc}", file=sys.stderr)
            return None
    except OSError as exc:
        print(f"[filesystem] ⚠️ Errore leggendo {path}: {exc}", file=sys.stderr)
        return None

    normalized = strip_front_matter(raw)
    cleaned = normalize_text(normalized)
    return cleaned or None


def read_pdf(path: Path) -> Optional[str]:
    """Legge testo da PDF con pymupdf; se il testo è assente tenta OCR."""

    try:
        import fitz  # pymupdf
    except ImportError:
        warn_missing_dependency("pymupdf", "lettura testi PDF")
        return _read_pdf_fallback(path)

    text_parts: List[str] = []
    try:
        doc = fitz.open(str(path))
        for page in doc:
            raw_text: str = page.get_text("text") or ""  # type: ignore[assignment]
            snippet = normalize_text(raw_text)
            if snippet:
                text_parts.append(snippet)
        doc.close()
    except Exception as exc:
        print(f"[filesystem] ⚠️ Errore leggendo PDF {path}: {exc}", file=sys.stderr)
        return None

    combined = "\n\n".join(part for part in text_parts if part)
    if combined.strip():
        return combined

    # Nessun testo estratto → documento scannerizzato, prova OCR
    return run_pdf_ocr(path)


def _read_pdf_fallback(path: Path) -> Optional[str]:
    """Fallback con pypdf se pymupdf non è disponibile."""
    try:
        from pypdf import PdfReader
    except ImportError:
        warn_missing_dependency("pypdf", "lettura testi PDF")
        return None

    text_parts: List[str] = []
    try:
        reader = PdfReader(str(path))
        for page in reader.pages:
            snippet = normalize_text(page.extract_text() or "")
            if snippet:
                text_parts.append(snippet)
    except Exception as exc:
        print(f"[filesystem] ⚠️ Errore leggendo PDF {path}: {exc}", file=sys.stderr)
        return None

    combined = "\n\n".join(part for part in text_parts if part)
    if combined.strip():
        return combined
    return run_pdf_ocr(path)


def read_image(path: Path) -> Optional[str]:
    """Legge testo da immagini tramite OCR."""

    return run_image_ocr(path)


def read_word(path: Path) -> Optional[str]:
    """Legge documenti Word (docx/doc)."""

    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".doc":
        return _read_doc_legacy(path)
    return None


def _read_docx(path: Path) -> Optional[str]:
    try:
        from docx import Document  # type: ignore[import]
    except ImportError:  # pragma: no cover - dipendenza opzionale
        warn_missing_dependency("python-docx", "lettura documenti Word")
        return None

    try:
        document = Document(str(path))
    except Exception as exc:
        print(f"[filesystem] ⚠️ Errore leggendo Word {path}: {exc}", file=sys.stderr)
        return None

    text_parts: List[str] = []
    for paragraph in document.paragraphs:
        snippet = normalize_text(paragraph.text)
        if snippet:
            text_parts.append(snippet)

    for table in document.tables:
        for row in table.rows:
            cells = [normalize_text(cell.text) for cell in row.cells if cell.text]
            row_text = " | ".join(filter(None, cells))
            if row_text:
                text_parts.append(row_text)

    # Estrai e OCR immagini embedded
    image_texts = _extract_docx_images(document, path)
    text_parts.extend(image_texts)

    combined = "\n\n".join(text_parts)
    return combined.strip() or None


_RASTER_IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".bmp",
    ".tiff",
    ".tif",
    ".webp",
}


def _extract_docx_images(document, path: Path) -> List[str]:
    """Estrae immagini raster embedded dal DOCX ed esegue OCR su ciascuna.
    Salta formati vettoriali (WMF, EMF) non adatti all'OCR.
    """
    results: List[str] = []
    try:
        image_rels = [
            rel for rel in document.part.rels.values() if "image" in rel.reltype.lower()
        ]
        ocr_index = 0
        for rel in image_rels:
            # Filtra formati vettoriali non OCRabili
            part_name = getattr(rel.target_part, "partname", "") or ""
            ext = Path(str(part_name)).suffix.lower()
            if ext and ext not in _RASTER_IMAGE_EXTENSIONS:
                continue
            try:
                image_bytes = rel.target_part.blob
            except Exception:
                continue
            ocr_text = run_image_bytes_ocr(image_bytes, "Immagine")
            if ocr_text:
                ocr_index += 1
                results.append(f"[Immagine {ocr_index}]\n{ocr_text}")
    except Exception as exc:
        print(
            f"[filesystem] ⚠️ Estrazione immagini da DOCX {path}: {exc}",
            file=sys.stderr,
        )
    return results


def _read_doc_legacy(path: Path) -> Optional[str]:
    try:
        import textract  # type: ignore[import]
    except ImportError:  # pragma: no cover - dipendenza opzionale
        warn_missing_dependency("textract", "lettura documenti Word .doc")
        return None

    try:
        raw_bytes = textract.process(str(path))
    except Exception as exc:
        print(f"[filesystem] ⚠️ Errore leggendo Word {path}: {exc}", file=sys.stderr)
        return None

    try:
        text = raw_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = raw_bytes.decode("latin-1", errors="ignore")
    return normalize_text(text) or None


def read_powerpoint(path: Path) -> Optional[str]:
    """Estrae testo da presentazioni PowerPoint."""

    try:
        from pptx import Presentation  # type: ignore[import]
    except ImportError:  # pragma: no cover - dipendenza opzionale
        warn_missing_dependency("python-pptx", "lettura presentazioni PowerPoint")
        return None

    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        print(
            f"[filesystem] ⚠️ Errore leggendo PowerPoint {path}: {exc}", file=sys.stderr
        )
        return None

    text_parts: List[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        slide_lines: List[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame:
                paragraphs: List[str] = []
                for paragraph in shape.text_frame.paragraphs:
                    snippet = normalize_text(paragraph.text)
                    if snippet:
                        paragraphs.append(snippet)
                if paragraphs:
                    slide_lines.append("\n".join(paragraphs))
                continue

            if getattr(shape, "has_table", False):
                table_lines: List[str] = []
                for row in shape.table.rows:  # type: ignore[union-attr]
                    cells: List[str] = []
                    for cell in row.cells:
                        snippet = normalize_text(cell.text)
                        if snippet:
                            cells.append(snippet)
                    if cells:
                        table_lines.append(" | ".join(cells))
                if table_lines:
                    slide_lines.append("\n".join(table_lines))
                continue

            text = getattr(shape, "text", "")
            snippet = normalize_text(text)
            if snippet:
                slide_lines.append(snippet)
        if slide_lines:
            slide_text = f"[Slide {index}]\n" + "\n".join(slide_lines)
            text_parts.append(slide_text)

    combined = "\n\n".join(text_parts)
    return combined.strip() or None


def read_excel(path: Path) -> Optional[str]:
    """Estrae dati leggibili da file Excel (xlsx/xls).

    Produce output strutturato: ogni riga di dati viene convertita in un record
    con intestazioni di colonna associate, così il testo resta comprensibile
    anche dopo il chunking per indicizzazione vettoriale.
    """

    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _read_excel_xlsx(path)
    if suffix == ".xls":
        return _read_excel_xls(path)
    return None


# ---------------------------------------------------------------------------
# Helpers per la conversione strutturata delle righe Excel
# ---------------------------------------------------------------------------

_MIN_HEADER_COLS = (
    3  # Soglia minima di colonne testuali per considerare una riga come header
)


def _cell_str(value: object) -> str:
    """Converte un valore di cella in stringa normalizzata (vuota se None/blank)."""
    if value is None:
        return ""
    text = normalize_text(str(value))
    return text if text else ""


def _text_cell_count(row_values: List[str]) -> int:
    """Conta le celle con valori testuali (non numerici/booleani)."""
    count = 0
    for v in row_values:
        if not v:
            continue
        stripped = v.replace(".", "").replace(",", "").replace("-", "")
        if stripped.isdigit():
            continue
        if v.lower() in ("true", "false", "vero", "falso"):
            continue
        count += 1
    return count


def _is_likely_header(row_values: List[str]) -> bool:
    """Euristica: una riga è intestazione se ha almeno *_MIN_HEADER_COLS* valori
    testuali non numerici."""
    return _text_cell_count(row_values) >= _MIN_HEADER_COLS


def _format_sheet_structured(
    rows: List[List[str]],
    sheet_name: str,
) -> Optional[str]:
    """Converte le righe di un foglio in testo strutturato.

    Strategia:
    1. Le righe prima della tabella dati con pochi valori (<=2 celle piene)
       vengono emesse come metadati del foglio (chiave: valore).
    2. La prima riga con >= _MIN_HEADER_COLS valori testuali diventa l'header.
    3. Ogni riga dati successiva diventa un record ``Header: Valore``.
    """
    if not rows:
        return None

    metadata_lines: List[str] = []
    header: Optional[List[str]] = None
    data_start = 0
    last_metadata_idx = -1

    # Fase 1: individua metadati e riga header
    for idx, row in enumerate(rows):
        non_empty = [v for v in row if v]
        if not non_empty:
            continue

        ne_count = len(non_empty)

        # Riga con 1-2 valori non vuoti prima dell'header → metadato chiave-valore
        if header is None and ne_count <= 2:
            if ne_count == 2:
                metadata_lines.append(f"{non_empty[0]}: {non_empty[1]}")
            else:
                metadata_lines.append(non_empty[0])
            last_metadata_idx = idx
            continue

        if header is None and _is_likely_header(row):
            header = row
            data_start = idx + 1
            break
        # Se non è header ma ha molti valori, inizio dati senza header esplicito
        if header is None and ne_count >= _MIN_HEADER_COLS:
            data_start = idx
            break

    # Se non abbiamo trovato header né dati, data_start parte dopo i metadati
    if header is None and data_start == 0 and last_metadata_idx >= 0:
        data_start = last_metadata_idx + 1

    parts: List[str] = []
    parts.append(f"[Foglio {sheet_name}]")
    if metadata_lines:
        parts.append("\n".join(metadata_lines))

    # Fase 2: emetti i record
    data_rows = rows[data_start:]
    if not data_rows:
        # Foglio con solo metadati/poche righe
        return "\n".join(parts).strip() or None

    # Ripeti il tag foglio ogni _SHEET_HEADER_REPEAT record, così il chunker
    # (350 char) lascia il contesto del foglio in ogni chunk risultante.
    _SHEET_HEADER_REPEAT = 5
    sheet_tag = f"[Foglio {sheet_name}]"
    record_count = 0

    if header is not None:
        for row in data_rows:
            fields: List[str] = []
            for col_idx, value in enumerate(row):
                if not value:
                    continue
                col_name = (
                    header[col_idx]
                    if col_idx < len(header) and header[col_idx]
                    else f"Col{col_idx + 1}"
                )
                fields.append(f"{col_name}: {value}")
            if fields:
                record_count += 1
                if record_count % _SHEET_HEADER_REPEAT == 0:
                    parts.append(sheet_tag)
                parts.append(" | ".join(fields))
    else:
        # Nessun header rilevato → output piatto come fallback
        for row in data_rows:
            non_empty = [v for v in row if v]
            if non_empty:
                record_count += 1
                if record_count % _SHEET_HEADER_REPEAT == 0:
                    parts.append(sheet_tag)
                parts.append(" | ".join(non_empty))

    # Doppio newline tra record: il chunker usa "\n\n" come separatore
    # primario, così ogni record resta intero dopo lo splitting.
    combined = "\n\n".join(parts)
    return combined.strip() or None


# ---------------------------------------------------------------------------
# Reader XLSX (openpyxl)
# ---------------------------------------------------------------------------


def _read_excel_xlsx(path: Path) -> Optional[str]:
    try:
        from openpyxl import load_workbook  # type: ignore[import]
    except ImportError:  # pragma: no cover - dipendenza opzionale
        warn_missing_dependency("openpyxl", "lettura fogli Excel xlsx")
        return None

    try:
        workbook = load_workbook(
            filename=str(path),
            data_only=True,
            read_only=True,
        )
    except Exception as exc:
        print(f"[filesystem] ⚠️ Errore leggendo Excel {path}: {exc}", file=sys.stderr)
        return None

    text_parts: List[str] = []
    try:
        for sheet in workbook.worksheets:
            all_rows: List[List[str]] = []
            for row in sheet.iter_rows(values_only=True):
                all_rows.append([_cell_str(v) for v in row])
            formatted = _format_sheet_structured(all_rows, sheet.title)
            if formatted:
                text_parts.append(formatted)
    finally:
        workbook.close()

    combined = "\n\n".join(text_parts)
    return combined.strip() or None


# ---------------------------------------------------------------------------
# Reader XLS (xlrd)
# ---------------------------------------------------------------------------


def _read_excel_xls(path: Path) -> Optional[str]:
    try:
        import xlrd  # type: ignore[import]
    except ImportError:  # pragma: no cover - dipendenza opzionale
        warn_missing_dependency("xlrd", "lettura fogli Excel xls")
        return None

    try:
        workbook = xlrd.open_workbook(str(path))
    except Exception as exc:
        print(f"[filesystem] ⚠️ Errore leggendo Excel {path}: {exc}", file=sys.stderr)
        return None

    text_parts: List[str] = []
    try:
        for sheet_index in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(sheet_index)
            all_rows: List[List[str]] = []
            for row_idx in range(sheet.nrows):
                row = [
                    _cell_str(sheet.cell_value(row_idx, col_idx))
                    for col_idx in range(sheet.ncols)
                ]
                all_rows.append(row)
            formatted = _format_sheet_structured(all_rows, sheet.name)
            if formatted:
                text_parts.append(formatted)
    finally:
        try:
            workbook.release_resources()
        except Exception:  # pragma: no cover - pulizia best effort
            pass

    combined = "\n\n".join(text_parts)
    return combined.strip() or None

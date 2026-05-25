"""XLSX parser via openpyxl."""

import uuid
from pathlib import Path

from ...schemas.state import Chunk


def parse(path: Path) -> list[Chunk]:
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(path), data_only=True, read_only=True)
    chunks: list[Chunk] = []
    for sheet in wb.worksheets:
        for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            text = " | ".join(str(c) for c in row if c is not None)
            if not text.strip():
                continue
            chunks.append(
                Chunk(
                    id=f"c-{uuid.uuid4().hex[:8]}",
                    text=f"[{sheet.title}] {text}",
                    page=1,
                    line_start=row_idx,
                    line_end=row_idx,
                    bbox=(0.0, 0.0, 0.0, 0.0),
                    token_count=len(text.split()),
                )
            )
    return chunks
